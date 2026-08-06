from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "static" / "downloads" / "intake-kit"
OUT.mkdir(parents=True, exist_ok=True)

INK = "1D1F20"
BRASS = "9B7438"
MUTED = "646A6E"
LINE = "C9C2B6"
PALE = "F3EFE6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_document(doc: Document, title: str, subject: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.24)

    normal = doc.styles["Normal"]
    normal.font.name = "Inter"
    normal.font.size = Pt(9.1)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 29, INK),
        ("Subtitle", 11, MUTED),
        ("Heading 1", 18, INK),
        ("Heading 2", 11.5, INK),
        ("Heading 3", 9.5, BRASS),
    ):
        style = doc.styles[style_name]
        style.font.name = "Inter"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(7 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(3)

    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "Forge & Verse"
    props.keywords = "digital forensics, intake, request, law enforcement, evidence"
    props.comments = "Forge & Verse free operational release. Adapt to local policy and approved workflows."

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(7.25))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.7)
    table.columns[1].width = Inches(2.55)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    for cell in (left, right):
        set_cell_margins(cell, 0, 0, 20, 0)
        set_cell_border(cell, bottom={"val": "single", "sz": "8", "color": BRASS})
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("DIGITAL FORENSICS")
    run.font.name = "Inter"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BRASS)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("FORGE & VERSE")
    run.font.name = "Inter"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Forge & Verse free release  |  Intake Kit v1.2  |  Adapt to agency policy, law, records requirements, and approved workflows.")
    run.font.name = "Inter"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Inter"
    r.font.size = Pt(29)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(subtitle)
    r.font.name = "Inter"
    r.font.size = Pt(10.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_section_bar(doc: Document, number: str, title: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.48)
    table.columns[1].width = Inches(6.77)
    c1, c2 = table.rows[0].cells
    set_cell_shading(c1, BRASS)
    set_cell_shading(c2, INK)
    for cell in (c1, c2):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 70, 95, 70, 95)
        set_cell_border(cell, top={"val": "nil"}, left={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(number)
    r.font.name = "Inter"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = c2.paragraphs[0]
    r = p.add_run(title.upper())
    r.font.name = "Inter"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: Iterable[str], size: float = 8.8, spacing: float = 2.5) -> None:
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.paragraph_format.space_after = Pt(spacing)
        p.paragraph_format.line_spacing = 1.04
        r = p.add_run(text)
        r.font.name = "Inter"
        r.font.size = Pt(size)


def add_numbered(doc: Document, items: Iterable[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.42)
    table.columns[1].width = Inches(6.83)
    for idx, (title, text) in enumerate(items, 1):
        cells = table.add_row().cells
        set_cell_margins(cells[0], 55, 20, 55, 20)
        set_cell_margins(cells[1], 55, 80, 55, 20)
        p = cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{idx:02}")
        r.font.name = "Inter"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BRASS)
        p = cells[1].paragraphs[0]
        r = p.add_run(title + "  ")
        r.font.name = "Inter"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
        r = p.add_run(text)
        r.font.name = "Inter"
        r.font.size = Pt(8.7)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        if idx < len(list(items)):
            set_cell_border(cells[0], bottom={"val": "single", "sz": "5", "color": LINE})
            set_cell_border(cells[1], bottom={"val": "single", "sz": "5", "color": LINE})


def add_callout(doc: Document, label: str, text: str, dark: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_margins(cell, 120, 150, 120, 150)
    set_cell_shading(cell, INK if dark else PALE)
    set_cell_border(cell, top={"val": "single", "sz": "8", "color": BRASS}, left={"val": "single", "sz": "4", "color": LINE}, bottom={"val": "single", "sz": "4", "color": LINE}, right={"val": "single", "sz": "4", "color": LINE})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    r.font.name = "Inter"
    r.font.size = Pt(7.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BRASS if dark else BRASS)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Inter"
    r.font.size = Pt(10 if dark else 9)
    r.font.bold = dark
    r.font.color.rgb = RGBColor.from_string(WHITE if dark else INK)


def add_two_column_table(doc: Document, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.58)
    table.columns[1].width = Inches(3.58)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, INK)
        set_cell_margins(cell, 80, 110, 80, 110)
    for idx, title in enumerate((left_title, right_title)):
        p = table.cell(0, idx).paragraphs[0]
        r = p.add_run(title.upper())
        r.font.name = "Inter"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
    for idx, items in enumerate((left_items, right_items)):
        cell = table.cell(1, idx)
        set_cell_margins(cell, 100, 120, 90, 120)
        set_cell_shading(cell, "FAF9F6")
        set_cell_border(cell, top={"val": "single", "sz": "4", "color": LINE}, left={"val": "single", "sz": "4", "color": LINE}, bottom={"val": "single", "sz": "4", "color": LINE}, right={"val": "single", "sz": "4", "color": LINE})
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
        for item in items:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.14)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            r = p.add_run("• ")
            r.font.name = "Inter"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(BRASS)
            r = p.add_run(item)
            r.font.name = "Inter"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(INK)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_requester_guide() -> Path:
    doc = Document()
    configure_document(doc, "Before You Submit a Digital Forensics Request", "One-page requester guide for officers and investigators")
    add_title(doc, "Before You Submit a Digital Forensics Request", "A one-page guide for officers and investigators")

    add_callout(doc, "Begin with the question", "The examiner needs to know what the investigation is trying to establish, locate, confirm, or exclude. A requested tool action is not the same thing as an investigative objective.", dark=True)

    add_section_bar(doc, "01", "Give the examiner a question they can act on")
    add_two_column_table(
        doc,
        "Useful objectives",
        [
            "Identify communications between named parties during a defined date range.",
            "Determine whether the device was used to access a named account or service.",
            "Locate images, documents, or location data relevant to a specific incident.",
        ],
        "Requests that need clarification",
        [
            '"Dump the phone."',
            '"Get everything."',
            '"Run Cellebrite on it."',
            "A list of tools without the question the evidence should help answer.",
        ],
    )

    add_section_bar(doc, "02", "Before the handoff")
    checklist = [
        ("Authority", "Identify the authority, where the documentation is stored, and any expiration or limits."),
        ("Scope", "Name the devices, accounts, data types, date ranges, and people covered by the request."),
        ("Targets", "Provide known contacts, usernames, phone numbers, aliases, applications, keywords, and dates."),
        ("Context", "Explain facts that may change how the request is approached, including remote-wipe risk or time-sensitive concerns."),
        ("Condition", "Record whether the device is on or off, networked or isolated, damaged, wet, or previously accessed."),
        ("Credentials", "State whether credentials are known, but transmit them only through the agency-approved method."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.42)
    table.columns[1].width = Inches(6.83)
    for idx, (title, text) in enumerate(checklist, 1):
        cells = table.add_row().cells
        set_cell_margins(cells[0], 45, 20, 45, 20)
        set_cell_margins(cells[1], 45, 80, 45, 20)
        p = cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[ ]")
        r.font.name = "Inter"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BRASS)
        p = cells[1].paragraphs[0]
        r = p.add_run(title + ": ")
        r.font.name = "Inter"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
        r = p.add_run(text)
        r.font.name = "Inter"
        r.font.size = Pt(8.3)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        if idx < len(checklist):
            set_cell_border(cells[0], bottom={"val": "single", "sz": "4", "color": LINE})
            set_cell_border(cells[1], bottom={"val": "single", "sz": "4", "color": LINE})

    add_section_bar(doc, "03", "Protect the evidence")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run("Follow agency policy for handling, isolation, packaging, and transfer. Do not explore the device, change settings, or place sensitive credentials on the intake form. Document any preservation action or change already made.")
    add_callout(doc, "A complete request does not replace a conversation", "When authority, scope, device condition, safety, or urgency is unclear, contact the forensic unit before taking additional action.")

    path = OUT / "digital-forensics-intake-requester-guide.docx"
    doc.save(path)
    return path


def build_implementation_guide() -> Path:
    doc = Document()
    configure_document(doc, "Digital Forensics Intake Kit Implementation Guide", "Guide for adapting and deploying the Forge & Verse intake release")
    add_title(doc, "Digital Forensics Intake Kit", "Implementation guide for forensic units and agency administrators")
    add_callout(doc, "Purpose", "This guide helps a unit adapt the intake form and requester guide to its actual authority, services, routing, evidence systems, records obligations, and approval structure. It is not a universal policy or legal form.", dark=True)

    add_section_bar(doc, "01", "Decide what the intake must accomplish")
    p = doc.add_paragraph("The form should create a reliable handoff between the person who understands the investigation and the person who must examine the evidence. It should reduce avoidable clarification without asking requesters to make forensic decisions they are not equipped to make.")
    p.paragraph_format.space_after = Pt(5)
    add_two_column_table(
        doc,
        "The intake should capture",
        [
            "Who is requesting the work and how to reach them.",
            "The authority, scope, limits, and location of supporting documentation.",
            "The question the examination should help answer.",
            "Submitted items, identifiers, condition, handling, and known access information.",
            "Targets, date ranges, known accounts, relevant context, and urgency.",
        ],
        "The intake should not attempt to",
        [
            "Replace legal review, local policy, or a warrant attachment.",
            "Guarantee that a particular tool or data source will produce a result.",
            "Ask requesters to select technical actions without explaining the investigative need.",
            "Collect passwords or sensitive credentials through an unapproved channel.",
            "Become so comprehensive that requesters bypass it and return to informal messages.",
        ],
    )
    add_callout(doc, "Design test", "Can the examiner understand the authority, the question, the evidence, the targets, and the immediate risks without reconstructing the request from several emails?")

    page_break(doc)
    add_title(doc, "Adapt the Template", "Change the form before it becomes part of the workflow")
    add_section_bar(doc, "02", "Local decisions to make")
    adaptation_rows = [
        ("Agency identity", "Add the unit name, routing address, contact information, revision owner, and approved branding."),
        ("Authority", "Use choices and terminology that match local law, policy, prosecutor expectations, and the work your unit accepts."),
        ("Services", "Remove services the unit does not offer. Add local categories only when requesters can understand and use them consistently."),
        ("Priority", "Define what routine, expedited, and emergency mean. Avoid allowing every requested deadline to become a priority by default."),
        ("Evidence systems", "Add item numbers, transfer records, queue numbers, or case-management identifiers needed to reconcile the request with custody records."),
        ("Approvals", "Add supervisor, prosecutor, unit, or budget approvals only when they are required and have a clear owner."),
        ("Signatures", "Decide whether signatures are required, whether electronic acknowledgment is accepted, and where the completed intake becomes a record."),
        ("Format and viewer", "Choose whether requesters use the print PDF, local fillable PDF, or an adapted agency form. Test saving, reopening, printing, and routing in the approved PDF viewer before publication."),
        ("Credentials", "State the approved channel for passcodes and account credentials. The form should never invite sensitive values into an unsecured workflow."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(5.8)
    headers = table.rows[0].cells
    for idx, label in enumerate(("Decision", "What to resolve")):
        set_cell_shading(headers[idx], INK)
        set_cell_margins(headers[idx], 75, 100, 75, 100)
        p = headers[idx].paragraphs[0]
        r = p.add_run(label.upper())
        r.font.name = "Inter"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(table.rows[0])
    for idx, (label, text) in enumerate(adaptation_rows):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, 70, 100, 70, 100)
            set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": LINE}, left={"val": "single", "sz": "4", "color": LINE}, right={"val": "single", "sz": "4", "color": LINE})
        if idx % 2 == 0:
            set_cell_shading(cells[0], "F8F5EE")
            set_cell_shading(cells[1], "F8F5EE")
        p = cells[0].paragraphs[0]
        r = p.add_run(label)
        r.font.name = "Inter"
        r.font.size = Pt(8.3)
        r.font.bold = True
        p = cells[1].paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Inter"
        r.font.size = Pt(8.2)
    add_callout(doc, "Keep the form short enough to use", "A field earns its place only when someone uses the answer to route, scope, preserve, prioritize, examine, document, or accept the request.")
    add_callout(doc, "The fillable edition stays local", "The Forge & Verse fillable PDF contains no automatic submission action, JavaScript, telemetry, or connection to the studio. The adopting agency still decides where completed files are stored, transmitted, retained, and protected.")

    page_break(doc)
    add_title(doc, "Roll It Out Deliberately", "A form succeeds when the surrounding handoff is clear")
    add_section_bar(doc, "03", "Suggested adoption path")
    steps = [
        ("Assign an owner", "Name the person or role responsible for the current version, approved changes, and the authoritative download location."),
        ("Review the workflow", "Include examiners, investigators, evidence personnel, supervisors, records staff, legal counsel, and any team that must act on the form."),
        ("Adapt and test", "Customize the files, then run realistic requests through the full routing and acceptance process. Test saving, reopening, printing, and records handling in the actual approved viewer."),
        ("Pilot with a small group", "Use a defined pilot period. Track missing information, unclear fields, workarounds, and repeated questions."),
        ("Teach the question", "Explain that the forensic unit needs the investigative objective, not merely a requested extraction or tool name."),
        ("Publish one source", "Place the approved version in one controlled location and remove older copies from shared drives, forms libraries, and email templates."),
        ("Review after use", "Revisit the form after the pilot and on a scheduled basis. Change it when the workflow changes, not whenever one unusual case appears."),
    ]
    for idx, (title, text) in enumerate(steps, 1):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(0.54)
        table.columns[1].width = Inches(6.71)
        number, body = table.rows[0].cells
        set_cell_shading(number, BRASS)
        set_cell_shading(body, "FAF9F6")
        set_cell_margins(number, 80, 30, 80, 30)
        set_cell_margins(body, 80, 110, 80, 110)
        p = number.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{idx:02}")
        r.font.name = "Inter"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        p = body.paragraphs[0]
        r = p.add_run(title + "  ")
        r.font.name = "Inter"
        r.font.size = Pt(8.7)
        r.font.bold = True
        r = p.add_run(text)
        r.font.name = "Inter"
        r.font.size = Pt(8.4)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_callout(doc, "Minimum training message", "Describe the question. Identify the authority and limits. Tell the examiner what was submitted, what happened to it, what is time-sensitive, and who can clarify the request.")

    page_break(doc)
    add_title(doc, "Field Rationale", "Why the form asks what it asks")
    add_section_bar(doc, "04", "High-value fields")
    rationale = [
        ("Question the examination should help answer", "Turns a tool request into an investigative objective and gives the examiner a basis for prioritization, searching, interpretation, and reporting."),
        ("Scope and express limitations", "Makes boundaries visible at intake rather than after examination has begun."),
        ("Known facts and time-sensitive concerns", "Surfaces remote-wipe risk, disappearing data, hearings, safety issues, or operational deadlines that may change the approach."),
        ("Device state and handling history", "Documents whether the item was powered, networked, accessed, isolated, damaged, wet, or changed before submission."),
        ("Known contacts, accounts, keywords, and dates", "Gives the examiner defensible targets instead of an undefined request to search everything."),
        ("Unit acceptance and clarification", "Creates a visible handoff where scope can be confirmed, missing information recorded, and the request assigned to the local queue or case system."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.32)
    table.columns[1].width = Inches(4.93)
    for idx, label in enumerate(("Field", "Why it matters")):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, INK)
        set_cell_margins(cell, 75, 100, 75, 100)
        r = cell.paragraphs[0].add_run(label.upper())
        r.font.name = "Inter"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(table.rows[0])
    for idx, (field, reason) in enumerate(rationale):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, 65, 100, 65, 100)
            set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": LINE}, left={"val": "single", "sz": "4", "color": LINE}, right={"val": "single", "sz": "4", "color": LINE})
        if idx % 2 == 0:
            set_cell_shading(cells[0], "F8F5EE")
            set_cell_shading(cells[1], "F8F5EE")
        r = cells[0].paragraphs[0].add_run(field)
        r.font.name = "Inter"
        r.font.size = Pt(8.2)
        r.font.bold = True
        r = cells[1].paragraphs[0].add_run(reason)
        r.font.name = "Inter"
        r.font.size = Pt(8.1)
    add_section_bar(doc, "05", "Version and records discipline")
    add_bullets(doc, [
        "Place a version number and effective date on the approved local form.",
        "Keep a brief change log describing what changed, who approved it, and when the new version became effective.",
        "State where completed requests are stored and how they relate to evidence, case, queue, and examination records.",
        "Remove superseded copies from common download locations and communicate the change to requesters.",
        "Review the form when authority, services, tools, routing, records obligations, or evidence systems change.",
    ], size=8.5, spacing=2)
    add_callout(doc, "Local review required", "Before operational use, review the adapted form and guide against applicable law, agency policy, records requirements, prosecutor expectations, evidence procedures, and the services the forensic unit actually provides.", dark=True)

    path = OUT / "digital-forensics-intake-implementation-guide.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    request = build_requester_guide()
    implementation = build_implementation_guide()
    print(request)
    print(implementation)
